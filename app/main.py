from __future__ import annotations

from decimal import Decimal, InvalidOperation
from io import BytesIO
from pathlib import Path
import re
from threading import Lock
from typing import Any

from fastapi import FastAPI, File, HTTPException, UploadFile, status
from openpyxl import load_workbook
from pydantic import BaseModel, Field

MIN_WEEKLY_LOAD = Decimal("0")
MAX_WEEKLY_LOAD = Decimal("5")
ALLOWED_EXTENSIONS = {".xlsx", ".xlsm"}

HEADER_ALIASES = {
    "surname": {
        "surname",
        "last_name",
        "lastname",
        "family_name",
        "familyname",
        "фамилия",
    },
    "weekly_load_days": {
        "weekly_load",
        "load_per_week",
        "week_load",
        "load_days",
        "weekly_load_days",
        "загрузка",
        "нагрузка",
        "загрузка_в_неделю",
        "нагрузка_в_неделю",
        "дней_в_неделю",
    },
}


class WorkloadItem(BaseModel):
    surname: str = Field(min_length=1, description="Employee surname")
    weekly_load_days: float = Field(
        ge=0,
        le=5,
        multiple_of=0.5,
        description="Weekly load in days from 0 to 5 with step 0.5",
    )


class UploadResponse(BaseModel):
    loaded_count: int
    items: list[WorkloadItem]


class WorkloadSummary(BaseModel):
    employee_count: int
    total_load_days: float
    average_load_days: float
    min_load_days: float
    max_load_days: float


class WorkloadStore:
    def __init__(self) -> None:
        self._data: dict[str, WorkloadItem] = {}
        self._lock = Lock()

    @staticmethod
    def _key(surname: str) -> str:
        return surname.strip().casefold()

    def replace_all(self, items: list[WorkloadItem]) -> None:
        with self._lock:
            self._data = {self._key(item.surname): item for item in items}

    def get_all(self) -> list[WorkloadItem]:
        with self._lock:
            values = [item.model_copy() for item in self._data.values()]
        return sorted(values, key=lambda item: item.surname.casefold())

    def get_one(self, surname: str) -> WorkloadItem | None:
        key = self._key(surname)
        with self._lock:
            item = self._data.get(key)
            return item.model_copy() if item else None


def normalize_header(raw: Any) -> str:
    text = str(raw).strip().lower().replace(".", "")
    text = re.sub(r"[\s-]+", "_", text)
    return re.sub(r"_+", "_", text).strip("_")


def parse_surname(raw: Any, row_number: int) -> str:
    if raw is None:
        raise ValueError(f"Row {row_number}: surname is required.")

    surname = str(raw).strip()
    if not surname:
        raise ValueError(f"Row {row_number}: surname is required.")
    return surname


def parse_weekly_load(raw: Any, row_number: int) -> float:
    if raw is None or str(raw).strip() == "":
        raise ValueError(f"Row {row_number}: weekly load is required.")

    try:
        value = Decimal(str(raw).strip().replace(",", "."))
    except InvalidOperation as exc:
        raise ValueError(f"Row {row_number}: weekly load must be numeric.") from exc

    if value < MIN_WEEKLY_LOAD or value > MAX_WEEKLY_LOAD:
        raise ValueError(f"Row {row_number}: weekly load must be in range [0, 5].")

    half_day_units = value * 2
    if half_day_units != half_day_units.to_integral_value():
        raise ValueError(f"Row {row_number}: weekly load step must be 0.5.")

    return float(value)


def detect_columns(header_row: tuple[Any, ...]) -> tuple[int, int]:
    surname_idx: int | None = None
    load_idx: int | None = None

    for idx, raw_header in enumerate(header_row):
        if raw_header is None:
            continue

        normalized = normalize_header(raw_header)
        if normalized in HEADER_ALIASES["surname"]:
            surname_idx = idx
            continue
        if normalized in HEADER_ALIASES["weekly_load_days"]:
            load_idx = idx

    if surname_idx is None or load_idx is None:
        expected_surname = ", ".join(sorted(HEADER_ALIASES["surname"]))
        expected_load = ", ".join(sorted(HEADER_ALIASES["weekly_load_days"]))
        raise ValueError(
            "Header row must contain surname and weekly load columns. "
            f"Surname aliases: {expected_surname}. "
            f"Weekly load aliases: {expected_load}."
        )

    return surname_idx, load_idx


def parse_excel_workloads(file_bytes: bytes) -> list[WorkloadItem]:
    workbook = load_workbook(filename=BytesIO(file_bytes), data_only=True)
    try:
        sheet = workbook.active

        header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), None)
        if not header_row:
            raise ValueError("Excel file is empty.")

        surname_idx, load_idx = detect_columns(header_row)

        items: list[WorkloadItem] = []
        seen_surnames: dict[str, int] = {}

        for row_number, row in enumerate(
            sheet.iter_rows(min_row=2, values_only=True),
            start=2,
        ):
            surname_raw = row[surname_idx] if surname_idx < len(row) else None
            load_raw = row[load_idx] if load_idx < len(row) else None

            if (
                (surname_raw is None or str(surname_raw).strip() == "")
                and (load_raw is None or str(load_raw).strip() == "")
            ):
                continue

            surname = parse_surname(surname_raw, row_number=row_number)
            weekly_load_days = parse_weekly_load(load_raw, row_number=row_number)

            duplicate_key = surname.casefold()
            if duplicate_key in seen_surnames:
                prev_row = seen_surnames[duplicate_key]
                raise ValueError(
                    f"Duplicate surname '{surname}' in row {row_number} "
                    f"(already defined in row {prev_row})."
                )
            seen_surnames[duplicate_key] = row_number

            items.append(
                WorkloadItem(surname=surname, weekly_load_days=weekly_load_days)
            )

        if not items:
            raise ValueError("No workload rows found after header.")

        return items
    finally:
        workbook.close()


def validate_filename(filename: str | None) -> None:
    if not filename:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Filename is missing.",
        )

    ext = Path(filename).suffix.lower()
    if ext not in ALLOWED_EXTENSIONS:
        allowed = ", ".join(sorted(ALLOWED_EXTENSIONS))
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Only Excel files are supported: {allowed}.",
        )


app = FastAPI(
    title="Weekly Workload Service",
    description=(
        "FastAPI service that loads employee surnames from an Excel file and "
        "stores weekly workload in days in the range 0..5 with step 0.5."
    ),
    version="1.0.0",
)

store = WorkloadStore()


@app.get("/health")
def healthcheck() -> dict[str, str]:
    return {"status": "ok"}


@app.post(
    "/api/v1/workloads/upload",
    response_model=UploadResponse,
    status_code=status.HTTP_201_CREATED,
)
async def upload_workloads(file: UploadFile = File(...)) -> UploadResponse:
    validate_filename(file.filename)
    content = await file.read()
    if not content:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Uploaded file is empty.",
        )

    try:
        items = parse_excel_workloads(content)
    except ValueError as exc:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail=str(exc),
        ) from exc
    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Failed to read Excel file.",
        ) from exc

    store.replace_all(items)
    saved_items = store.get_all()
    return UploadResponse(loaded_count=len(saved_items), items=saved_items)


@app.get("/api/v1/workloads", response_model=list[WorkloadItem])
def list_workloads() -> list[WorkloadItem]:
    return store.get_all()


@app.get("/api/v1/workloads/summary", response_model=WorkloadSummary)
def workload_summary() -> WorkloadSummary:
    items = store.get_all()
    if not items:
        return WorkloadSummary(
            employee_count=0,
            total_load_days=0.0,
            average_load_days=0.0,
            min_load_days=0.0,
            max_load_days=0.0,
        )

    loads = [item.weekly_load_days for item in items]
    total = sum(loads)
    average = total / len(loads)
    return WorkloadSummary(
        employee_count=len(items),
        total_load_days=round(total, 2),
        average_load_days=round(average, 2),
        min_load_days=min(loads),
        max_load_days=max(loads),
    )


@app.get("/api/v1/workloads/{surname}", response_model=WorkloadItem)
def get_workload_by_surname(surname: str) -> WorkloadItem:
    item = store.get_one(surname)
    if item is None:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail=f"Surname '{surname}' not found.",
        )
    return item
